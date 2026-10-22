"""The study's hand-kept session tracking workbook."""
from __future__ import annotations

import math
import re
import warnings
from dataclasses import dataclass, field
from datetime import datetime, timedelta
from pathlib import Path
from typing import Optional

import numpy as np
import openpyxl
import pandas as pd

_NOT_SUBJECTS = {"summary", "template", "text", "legend", "notes"}
_EXCEL_EPOCH = datetime(1899, 12, 30)

SESSION_COLUMNS = ["Subject", "Session", "Date", "Weekday", "OnMed", "MedTiming", "TaskOrder",
                   "Valid", "Note", "TasksDone", "RestDone"]


@dataclass
class Tracksheet:
    """What the tracking workbook says.

    Attributes
    ----------
    participants : DataFrame
        ``Subject, Group, Box, StartDate, EndDate`` from the Summary sheet, or
        ``Subject, Group`` (all ``'unassigned'``) when there is none.
    sessions : DataFrame
        One row per attempt: ``Subject, Session, Date, Weekday, OnMed,
        MedTiming, TaskOrder, Valid, Note, TasksDone, RestDone``.
    file : Path
    """
    participants: pd.DataFrame = field(default_factory=pd.DataFrame)
    sessions: pd.DataFrame = field(default_factory=pd.DataFrame)
    file: Optional[Path] = None


def read_tracksheet(path: str | Path) -> Tracksheet:
    """Read the session tracking workbook.

    The tracksheet is the study's own record of what was collected, kept by
    hand as sessions happened: which arm each subject is in, which attempt
    was the real session, what medication state it was in, and the order the
    tasks were run that day.

    Layout: a ``Summary`` sheet (Subject, Group, Box, Start date, End date),
    then one sheet per subject named with the id, one row per attempt::

        Session  Date  Weekday  On/off Med  Before/after Med
        Task 1 [status | name]  Task 2 [...]  Task 3 [...]  Rest  EMA  Comments
        Portal report [box | notes]  Tablet backup [box | notes]  ...

    Each "Task n" header sits over a 0/1 status box with the task name in the
    cell to its right. ``TasksDone`` counts the ticked boxes and is NaN when
    the row leaves all of them blank, which is not the same claim as a
    recorded zero. ``Note`` is the cell after "Tablet backup", where the
    reason for an excluded attempt is written.

    A row with no session number is an attempt that was not kept. Such rows
    are returned with ``Valid = False`` rather than dropped, so a scan can
    report what it is ignoring. Wholly blank rows are dropped.
    """
    path = Path(path)
    if not path.is_file():
        raise FileNotFoundError(f"No such file: {path}")
    try:
        with warnings.catch_warnings():
            # openpyxl warns about data-validation extensions it does not keep.
            warnings.simplefilter("ignore", UserWarning)
            book = openpyxl.load_workbook(path, read_only=True, data_only=True)
    except Exception as err:
        raise ValueError(f"{path} could not be opened as a workbook ({err}).") from err

    sheets = {name: [list(r) for r in book[name].iter_rows(values_only=True)]
              for name in book.sheetnames}
    book.close()

    out = Tracksheet(file=path)
    summary = next((n for n in sheets if n.lower() == "summary"), None)
    if summary is not None:
        out.participants = _read_summary(sheets[summary])

    frames = []
    for name, cells in sheets.items():
        if name.lower() in _NOT_SUBJECTS:
            continue
        frame = _read_subject_sheet(name.strip(), cells)
        if not frame.empty:
            frames.append(frame)
    out.sessions = (pd.concat(frames, ignore_index=True) if frames
                    else pd.DataFrame(columns=SESSION_COLUMNS))

    if out.participants.empty and not out.sessions.empty:
        subjects = pd.unique(out.sessions.Subject)
        out.participants = pd.DataFrame({"Subject": subjects, "Group": "unassigned"})
    return out


def _read_summary(cells: list[list]) -> pd.DataFrame:
    if not cells:
        return pd.DataFrame()
    hdr = _header(cells)
    c_subj = _find(hdr, "subject", "subject id", "id")
    if c_subj is None:
        return pd.DataFrame()
    c_group, c_box = _find(hdr, "group", "arm", "cohort"), _find(hdr, "box")
    c_start, c_end = _find(hdr, "start date", "start"), _find(hdr, "end date", "end")

    rows = []
    for r in cells[1:]:
        subject = _cell_str(_at(r, c_subj))
        if not subject:
            continue
        rows.append({"Subject": subject,
                     "Group": _cell_str(_at(r, c_group)) or "unassigned",
                     "Box": _cell_str(_at(r, c_box)),
                     "StartDate": _cell_date(_at(r, c_start)),
                     "EndDate": _cell_date(_at(r, c_end))})
    return pd.DataFrame(rows, columns=["Subject", "Group", "Box", "StartDate", "EndDate"])


def _read_subject_sheet(subject: str, cells: list[list]) -> pd.DataFrame:
    if len(cells) < 2:
        return pd.DataFrame()
    hdr = _header(cells)
    c_ses, c_date = _find(hdr, "session"), _find(hdr, "date")
    if c_ses is None or c_date is None:
        return pd.DataFrame()
    c_wd, c_rest = _find(hdr, "weekday"), _find(hdr, "rest")
    c_med = _find(hdr, "on/off med", "on off med", "med")
    c_tim = _find(hdr, "before/after med", "before after med")
    c_note = _find(hdr, "tablet backup")
    c_note = None if c_note is None else c_note + 1
    task_cols = [j for j, h in enumerate(hdr) if re.match(r"^task\s*\d", h)]
    name_cols = [j + 1 for j in task_cols if j + 1 < len(hdr)]

    rows = []
    for r in cells[1:]:
        session, date = _cell_str(_at(r, c_ses)), _cell_date(_at(r, c_date))
        if not session and date is None:
            continue
        flags = np.array([_cell_num(_at(r, j)) for j in task_cols])
        rows.append({
            "Subject": subject, "Session": session, "Date": date,
            "Weekday": _cell_str(_at(r, c_wd)),
            "OnMed": _cell_str(_at(r, c_med)).upper(),
            "MedTiming": _cell_str(_at(r, c_tim)),
            "TaskOrder": ">".join(v for j in name_cols if (v := _cell_str(_at(r, j)))),
            # A session number is what makes an attempt count.
            "Valid": bool(session) and date is not None,
            "Note": _cell_str(_at(r, c_note)),
            "TasksDone": float((flags == 1).sum()) if np.isfinite(flags).any() else math.nan,
            "RestDone": _cell_num(_at(r, c_rest)),
        })
    return pd.DataFrame(rows, columns=SESSION_COLUMNS)


def _header(cells: list[list]) -> list[str]:
    return [_cell_str(v).lower() for v in cells[0]]


def _find(hdr: list[str], *names: str) -> Optional[int]:
    for name in names:
        hit = next((j for j, h in enumerate(hdr) if h.startswith(name)), None)
        if hit is not None:
            return hit
    return None


def _at(row: list, col: Optional[int]) -> object:
    return row[col] if col is not None and col < len(row) else None


def _cell_str(v: object) -> str:
    if v is None or (isinstance(v, float) and math.isnan(v)):
        return ""
    if isinstance(v, float) and v.is_integer():
        v = int(v)
    if isinstance(v, bool):
        return "true" if v else "false"
    s = str(v).strip()
    return "" if s == "NaN" else s


def _cell_num(v: object) -> float:
    """A status box as a number; blank stays NaN rather than becoming a recorded zero."""
    if v is None:
        return math.nan
    if isinstance(v, (bool, int, float)):
        return float(v)
    try:
        return float(str(v).strip())
    except ValueError:
        return math.nan


def _cell_date(v: object) -> Optional[datetime]:
    """An Excel date, serial number or date written as text."""
    if v is None:
        return None
    if isinstance(v, datetime):
        return v
    if isinstance(v, (int, float)) and not isinstance(v, bool):
        # Below 20000 is 1954: far more likely a stray number than a study date.
        return _EXCEL_EPOCH + timedelta(days=float(v)) if 20000 < v < 80000 else None
    try:
        stamp = pd.Timestamp(str(v).strip())
    except ValueError:
        return None
    return None if pd.isna(stamp) else stamp.to_pydatetime()
